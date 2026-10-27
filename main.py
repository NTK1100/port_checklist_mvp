import os
import io
import json
import base64
import shutil
from datetime import datetime
from typing import List, Optional, Dict, Any

import fitz  # PyMuPDF
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font as XlFont, PatternFill, Alignment, Border, Side
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends, Response, BackgroundTasks, status
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from auth import (
    COOKIE_NAME,
    TOKEN_TTL_SECONDS,
    hash_password,
    verify_password,
    create_session_token,
    validate_file_security,
    sanitize_filename,
    get_current_user,
    require_roles,
)
from database import (
    init_db,
    get_connection,
    create_vessel,
    update_vessel_info,
    get_vessel_detail,
    get_all_vessels,
    recalculate_vessel_status,
    get_next_sequence_number,
    get_vessel_storage_abs_path,
    get_vessel_folder_meta_internal,
    get_file_record_internal,
    get_user_by_username,
    get_user_by_id,
    list_all_users,
    log_audit_event,
    get_audit_logs,
    VESSEL_TYPES,
    DOC_TYPE_INFO,
    STORAGE_ROOT,
    DB_PATH,
)
from parser_engine import parse_and_validate_file, build_standardized_filename, sanitize_token
from generate_samples import seed_initial_vessels, generate_all_sample_files, SAMPLE_DIR, create_pdf_file
from gdrive_bridge_client import (
    is_bridge_enabled,
    get_bridge_status,
    set_bridge_config,
    ping_bridge,
    load_db_from_gdrive,
    sync_db_to_gdrive,
    bg_create_vessel_folder_and_sync,
    bg_rename_vessel_folder_and_sync,
    bg_upload_doc_and_sync,
    bg_rename_doc_and_sync,
    bg_delete_doc_and_sync,
    bg_batch_delete_docs_and_sync,
    bg_clear_vessel_folder_and_sync,
    bg_delete_vessel_folder_and_sync,
    fetch_vessel_folder_from_gdrive,
    rename_doc_on_gdrive_sync,
    download_doc_from_gdrive_if_missing,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
FONTS_DIR = os.path.join(BASE_DIR, "fonts")

os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(STORAGE_ROOT, exist_ok=True)
os.makedirs(SAMPLE_DIR, exist_ok=True)
os.makedirs(FONTS_DIR, exist_ok=True)

app = FastAPI(
    title="PortCall Docs",
    version="2.5.0",
)


# ==================== PYDANTIC SCHEMAS ====================

class LoginRequest(BaseModel):
    username: str
    password: str


class VesselCreateItem(BaseModel):
    vessel_name: str
    eta: str
    vessel_type: str
    voyage_no: Optional[str] = ""


class VesselBatchCreateRequest(BaseModel):
    vessels: List[VesselCreateItem]


class VesselUpdateRequest(BaseModel):
    vessel_name: str
    eta: str
    vessel_type: Optional[str] = None


class ClassifyFileRequest(BaseModel):
    doc_type: str
    confirm_override: Optional[bool] = True
    replace_existing: Optional[bool] = False
    mark_checklist_completed: Optional[bool] = False


class ChecklistVerifyRequest(BaseModel):
    verified: bool = True
    doc_type_override: Optional[str] = None


class BatchDeleteFilesRequest(BaseModel):
    file_ids: List[int]


class BatchConfirmFilesRequest(BaseModel):
    file_ids: List[int]
    doc_type: Optional[str] = None


class SimulateSampleRequest(BaseModel):
    filenames: List[str]
    replace_existing: Optional[bool] = False


class UserCreateRequest(BaseModel):
    username: str
    full_name: str
    password: str
    role: str = "OPS"


class UserUpdateRequest(BaseModel):
    role: Optional[str] = None
    is_active: Optional[bool] = None
    full_name: Optional[str] = None
    password: Optional[str] = None


class GDriveBridgeConfigRequest(BaseModel):
    bridge_url: str
    secret_key: Optional[str] = None


class GDriveSyncDirectionRequest(BaseModel):
    direction: Optional[str] = "PUSH"


def _mime_for_ext(ext: str) -> str:
    ext = (ext or "").lower()
    if ext == ".pdf":
        return "application/pdf"
    if ext == ".docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if ext == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return "application/octet-stream"


def _resolve_vietnamese_fonts() -> Dict[str, Optional[str]]:
    """
    Tìm font Unicode hỗ trợ đầy đủ dấu Tiếng Việt (Regular & Bold) cho xuất PDF:
    1. Thư mục ./fonts/ đi kèm trong repository (VietnameseSans-Regular.ttf / VietnameseSans-Bold.ttf)
    2. Font DejaVuSans trên hệ điều hành Debian/Ubuntu Linux (Render Docker)
    3. Font Arial trên Windows
    """
    reg_candidates = [
        os.path.join(FONTS_DIR, "VietnameseSans-Regular.ttf"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
    ]
    bold_candidates = [
        os.path.join(FONTS_DIR, "VietnameseSans-Bold.ttf"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/segoeuib.ttf",
    ]
    reg_path = next((p for p in reg_candidates if os.path.exists(p)), None)
    bold_path = next((p for p in bold_candidates if os.path.exists(p)), reg_path)
    return {"regular": reg_path, "bold": bold_path}


# ==================== STARTUP & STATIC ====================

@app.on_event("startup")
def on_startup() -> None:
    """
    Khi server khởi động:
    1. Nếu GDRIVE_BRIDGE_URL được cấu hình, tự động gọi LOAD_DB để khôi phục port_checklist.db từ thư mục ẩn trên Google Drive.
    2. Khởi tạo/migrate bảng DB và sinh bộ file mẫu.
    3. Nếu trên Google Drive chưa có bản backup DB, tự động tạo lại từ dữ liệu hiện hành và đẩy lên Drive.
    """
    restored_from_cloud = False
    if is_bridge_enabled():
        restored_from_cloud = load_db_from_gdrive(DB_PATH)

    init_db()
    seed_initial_vessels(reset_existing=False)
    generate_all_sample_files()

    if is_bridge_enabled() and not restored_from_cloud:
        sync_db_to_gdrive(DB_PATH)


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/", include_in_schema=False)
def serve_spa():
    index_path = os.path.join(STATIC_DIR, "index.html")
    if not os.path.exists(index_path):
        return JSONResponse({"message": "Frontend index.html is being prepared."})
    return FileResponse(
        index_path,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )


# ==================== AUTHENTICATION API ====================

@app.post("/api/auth/login")
def login(payload: LoginRequest, response: Response):
    username = payload.username.strip()
    user = get_user_by_username(username)
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Tên đăng nhập hoặc mật khẩu không chính xác.",
        )

    if not user.get("is_active"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Tài khoản của bạn đã bị khóa bởi Quản trị viên.",
        )

    token = create_session_token(user["username"], user["role"])
    response.set_cookie(
        key=COOKIE_NAME,
        value=token,
        httponly=True,
        samesite="lax",
        max_age=TOKEN_TTL_SECONDS,
        path="/",
    )

    log_audit_event(
        username=user["username"],
        role=user["role"],
        action="LOGIN",
        details=f"Đăng nhập hệ thống ({user['full_name']})",
    )

    return {
        "message": "Đăng nhập thành công.",
        "user": {
            "id": user["id"],
            "username": user["username"],
            "full_name": user["full_name"],
            "role": user["role"],
            "is_active": bool(user["is_active"]),
        },
    }


@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie(key=COOKIE_NAME, path="/")
    return {"message": "Đã đăng xuất."}


@app.get("/api/auth/me")
def get_me(current_user: Dict[str, Any] = Depends(get_current_user)):
    return {"user": current_user}


# ==================== USER MANAGEMENT & AUDIT LOGS (ADMIN ONLY) ====================

@app.get("/api/users")
def api_list_users(current_user: Dict[str, Any] = Depends(require_roles("ADMIN"))):
    return {"users": list_all_users()}


@app.post("/api/users")
def api_create_user(
    payload: UserCreateRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN")),
):
    uname = payload.username.strip().lower()
    if not uname or len(uname) < 3:
        raise HTTPException(status_code=400, detail="Username phải có ít nhất 3 ký tự.")
    if payload.role not in ("ADMIN", "OPS", "VIEWER"):
        raise HTTPException(status_code=400, detail="Role không hợp lệ (ADMIN, OPS, VIEWER).")
    if not payload.password or len(payload.password) < 6:
        raise HTTPException(status_code=400, detail="Mật khẩu phải có ít nhất 6 ký tự.")
    if get_user_by_username(uname):
        raise HTTPException(status_code=400, detail=f"Tài khoản '{uname}' đã tồn tại.")

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO users (username, full_name, password_hash, role, is_active, created_at)
        VALUES (?, ?, ?, ?, 1, ?)
        """,
        (uname, payload.full_name.strip() or uname, hash_password(payload.password), payload.role, now_str),
    )
    conn.commit()
    conn.close()

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="USER_MGMT",
        details=f"Tạo tài khoản mới '{uname}' với quyền {payload.role}",
    )
    background_tasks.add_task(sync_db_to_gdrive)

    return {"message": f"Đã tạo tài khoản '{uname}'.", "users": list_all_users()}


@app.put("/api/users/{user_id}")
def api_update_user(
    user_id: int,
    payload: UserUpdateRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN")),
):
    target = get_user_by_id(user_id)
    if not target:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")

    if target["id"] == current_user["id"]:
        if payload.is_active is False:
            raise HTTPException(status_code=400, detail="Không thể tự khóa tài khoản đang đăng nhập.")
        if payload.role and payload.role != "ADMIN":
            raise HTTPException(status_code=400, detail="Không thể tự hạ quyền Quản trị viên của chính mình.")

    new_role = payload.role if payload.role in ("ADMIN", "OPS", "VIEWER") else target["role"]
    new_active = (1 if payload.is_active else 0) if payload.is_active is not None else target["is_active"]
    new_fullname = payload.full_name.strip() if payload.full_name else target["full_name"]
    new_pw_hash = hash_password(payload.password) if (payload.password and len(payload.password) >= 6) else target["password_hash"]

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        UPDATE users
        SET role = ?, is_active = ?, full_name = ?, password_hash = ?
        WHERE id = ?
        """,
        (new_role, new_active, new_fullname, new_pw_hash, user_id),
    )
    conn.commit()
    conn.close()

    changes = []
    if new_role != target["role"]:
        changes.append(f"Role: {target['role']} -> {new_role}")
    if new_active != target["is_active"]:
        changes.append("Trạng thái: " + ("Mở khóa" if new_active == 1 else "Khóa tài khoản"))
    if payload.password:
        changes.append("Đặt lại mật khẩu")

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="USER_MGMT",
        details=f"Cập nhật tài khoản '{target['username']}' ({', '.join(changes) or 'Cập nhật thông tin'})",
    )
    background_tasks.add_task(sync_db_to_gdrive)

    return {"message": f"Đã cập nhật tài khoản '{target['username']}'.", "users": list_all_users()}


@app.get("/api/audit-logs")
def api_get_audit_logs(current_user: Dict[str, Any] = Depends(require_roles("ADMIN"))):
    return {"logs": get_audit_logs(limit=250)}


# ==================== GOOGLE DRIVE BRIDGE ADMIN API ====================

@app.get("/api/gdrive/status")
def api_get_gdrive_status(current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER"))):
    return get_bridge_status()


@app.post("/api/admin/gdrive-config")
def api_save_gdrive_config(
    payload: GDriveBridgeConfigRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN")),
):
    set_bridge_config(payload.bridge_url, payload.secret_key)
    if not is_bridge_enabled():
        return {
            "message": "Đã chuyển về chế độ lưu trữ cục bộ (Local Storage).",
            "gdrive": get_bridge_status(),
        }

    ping_res = ping_bridge()
    if not ping_res.get("ok"):
        raise HTTPException(
            status_code=400,
            detail=f"Không thể kết nối tới Google Apps Script Bridge: {ping_res.get('error')}",
        )

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="USER_MGMT",
        details="Kết nối và kích hoạt Google Drive Bridge thành công",
    )
    background_tasks.add_task(sync_db_to_gdrive)

    return {
        "message": "Đã kết nối Google Drive Bridge và kích hoạt đồng bộ 2 chiều!",
        "gdrive": get_bridge_status(),
        "ping": {"status": "ok", **ping_res},
    }


@app.post("/api/admin/gdrive-sync")
def api_manual_gdrive_sync(
    payload: Optional[GDriveSyncDirectionRequest] = None,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    if not is_bridge_enabled():
        raise HTTPException(status_code=400, detail="Chưa cấu hình GDRIVE_BRIDGE_URL.")
    direction = (payload.direction if payload and payload.direction else "PUSH").upper()
    if direction == "PULL":
        ok = load_db_from_gdrive(DB_PATH)
        if not ok:
            sync_db_to_gdrive(DB_PATH)
            return {
                "message": "Chưa có bản backup trên Drive; đã tự động tạo bản backup mới từ dữ liệu hiện hành.",
                "gdrive": get_bridge_status(),
            }
        init_db()
        return {
            "message": "Đã khôi phục Database từ Google Drive thành công.",
            "gdrive": get_bridge_status(),
        }

    res = sync_db_to_gdrive()
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=f"Đồng bộ thất bại: {res.get('error')}")
    return {
        "message": "Đã đồng bộ Database lên thư mục ẩn Google Drive thành công.",
        "gdrive": get_bridge_status(),
    }


# ==================== INTERNAL FILE PROCESSING HELPER ====================

def _process_and_store_single_file(
    vessel: Dict[str, Any],
    temp_file_path: str,
    original_filename: str,
    uploaded_by: str,
    replace_existing: bool = False,
    preset_gdrive_file_id: str = "",
    preset_gdrive_view_url: str = "",
) -> Dict[str, Any]:
    """
    Đọc nội dung chứng từ, tự động phân loại & kiểm tra chéo tên tàu:
    - Nếu file bị cảnh báo WRONG_VESSEL hoặc UNCLASSIFIED:
      * Giữ ở trạng thái STAGED_WARNING (tuyệt đối KHÔNG tự động đẩy lên Google Drive).
    - Nếu file hợp lệ:
      * Đặt tên theo quy tắc [{LOAI_GIAY_TO}]_{Ten_Tau}.{ext}.
      * Nếu khác định dạng (.pdf vs .docx vs .xlsx) thì KHÔNG gắn _(2).
      * Nếu trùng loại chứng từ VÀ trùng định dạng:
        - Nếu replace_existing=True: Xóa bản cũ cùng định dạng và ghi đè tên gốc.
        - Nếu replace_existing=False: Lưu thành _(2) và gắn cờ has_same_ext_conflict=True để UI hỏi Ghi đè hay Giữ bổ sung.
    """
    vessel_id = vessel["id"]
    vessel_name = vessel["vessel_name"]

    abs_vessel_folder = get_vessel_storage_abs_path(vessel_id)
    if not abs_vessel_folder:
        raise HTTPException(status_code=500, detail="Không thể khởi tạo vùng lưu trữ cho tàu.")

    ext = os.path.splitext(original_filename)[1].lower() or ".pdf"

    parsed = parse_and_validate_file(
        file_path=temp_file_path,
        original_filename=original_filename,
        current_vessel_name=vessel_name,
    )
    doc_type = parsed["doc_type"]
    is_warning = bool(parsed["is_wrong_vessel"] == 1 or doc_type == "UNCLASSIFIED")

    conn = get_connection()
    cursor = conn.cursor()

    replaced_gdrive_items: List[Dict[str, str]] = []
    has_same_ext_conflict = False
    existing_same_ext_files: List[Dict[str, Any]] = []

    if not is_warning:
        cursor.execute(
            """
            SELECT id, stored_filename, file_path, gdrive_file_id
            FROM uploaded_files
            WHERE vessel_id = ? AND doc_type = ? AND LOWER(file_ext) = ? AND is_wrong_vessel = 0
            ORDER BY id ASC
            """,
            (vessel_id, doc_type, ext),
        )
        existing_same_ext_files = [dict(r) for r in cursor.fetchall()]

        if existing_same_ext_files and replace_existing:
            folder_meta = get_vessel_folder_meta_internal(vessel_id)
            for old_f in existing_same_ext_files:
                old_abs = os.path.join(abs_vessel_folder, os.path.basename(old_f["stored_filename"]))
                if os.path.exists(old_abs) and os.path.abspath(old_abs) != os.path.abspath(temp_file_path):
                    try:
                        os.remove(old_abs)
                    except Exception:
                        pass
                cursor.execute("DELETE FROM uploaded_files WHERE id = ?", (old_f["id"],))
                replaced_gdrive_items.append(
                    {
                        "gdrive_file_id": old_f.get("gdrive_file_id") or "",
                        "vessel_folder_name": folder_meta["folder_name"],
                        "stored_filename": old_f["stored_filename"],
                    }
                )
            seq_num = 1
        else:
            if existing_same_ext_files:
                has_same_ext_conflict = True
            seq_num = get_next_sequence_number(vessel_id, doc_type, file_ext=ext, conn=conn)

        stored_filename = build_standardized_filename(doc_type, vessel_name, seq_num, ext)
    else:
        # File đang ở trạng thái cảnh báo chờ User xác nhận -> Chưa đẩy lên Drive
        seq_num = get_next_sequence_number(vessel_id, doc_type, file_ext=ext, conn=conn)
        stored_filename = build_standardized_filename(doc_type, vessel_name, seq_num, ext)

    final_abs_path = os.path.join(abs_vessel_folder, stored_filename)
    while os.path.exists(final_abs_path) and os.path.abspath(final_abs_path) != os.path.abspath(temp_file_path):
        seq_num += 1
        stored_filename = build_standardized_filename(doc_type, vessel_name, seq_num, ext)
        final_abs_path = os.path.join(abs_vessel_folder, stored_filename)

    if os.path.abspath(temp_file_path) != os.path.abspath(final_abs_path):
        shutil.move(temp_file_path, final_abs_path)

    file_size = os.path.getsize(final_abs_path) if os.path.exists(final_abs_path) else 0
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if is_warning:
        initial_sync_status = "STAGED_WARNING"
    elif preset_gdrive_file_id:
        initial_sync_status = "SYNCED"
    else:
        initial_sync_status = "PENDING_SYNC" if is_bridge_enabled() else "LOCAL"

    cursor.execute(
        """
        INSERT INTO uploaded_files (
            vessel_id, original_filename, stored_filename, file_path,
            file_ext, file_size, doc_type, detected_method, detected_keyword,
            detected_vessel_name, is_wrong_vessel, error_code, error_message,
            extracted_count, extracted_codes_json, extracted_text_preview,
            uploaded_by, uploaded_at, gdrive_file_id, gdrive_view_url, sync_status
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            vessel_id,
            original_filename,
            stored_filename,
            stored_filename,
            ext,
            file_size,
            doc_type,
            parsed["detected_method"],
            parsed["detected_keyword"],
            parsed["detected_vessel_name"],
            parsed["is_wrong_vessel"],
            parsed["error_code"],
            parsed["error_message"],
            parsed["extracted_count"],
            json.dumps(parsed["extracted_codes"], ensure_ascii=False),
            parsed["extracted_text_preview"],
            uploaded_by,
            now_str,
            preset_gdrive_file_id,
            preset_gdrive_view_url,
            initial_sync_status,
        ),
    )
    file_id = cursor.lastrowid
    conn.commit()
    conn.close()

    return {
        "id": file_id,
        "original_filename": original_filename,
        "stored_filename": stored_filename,
        "abs_path": final_abs_path,
        "ext": ext,
        "doc_type": doc_type,
        "is_wrong_vessel": parsed["is_wrong_vessel"],
        "is_warning": is_warning,
        "has_same_ext_conflict": has_same_ext_conflict,
        "base_standard_filename": build_standardized_filename(doc_type, vessel_name, 1, ext) if not is_warning else "",
        "replaced_gdrive_items": replaced_gdrive_items,
        "error_message": parsed["error_message"],
        "extracted_count": parsed["extracted_count"],
        "extracted_codes": parsed["extracted_codes"],
        "uploaded_by": uploaded_by,
        "uploaded_at": now_str,
        "sync_status": initial_sync_status,
    }


# ==================== VESSELS & FILES API (PROTECTED BY RBAC) ====================

@app.get("/api/config")
def get_system_config(current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER"))):
    sample_descriptions = {
        "scan_001.pdf": "Cargo Manifest · Khai báo 5 B/L (BL-101..105)",
        "file_gui_sep.xlsx": "VGM Declaration (Excel) · 3/5 B/L (Thiếu 2 bộ)",
        "vgm_bo_sung.pdf": "VGM Declaration (PDF) · 2/5 B/L",
        "doc_bl_full.pdf": "Bill of Lading · Đủ 5/5 B/L",
        "crew_scan.pdf": "Crew List (PDF)",
        "crew_list_word.docx": "Crew List (Word .docx)",
        "port_clearance_99.pdf": "Port Clearance · Giấy phép rời cảng",
        "file_nham_tau.pdf": "VGM ghi sai tên tàu (MV EVER GIVEN)",
    }
    samples = []
    if os.path.exists(SAMPLE_DIR):
        for fname in sorted(os.listdir(SAMPLE_DIR)):
            fpath = os.path.join(SAMPLE_DIR, fname)
            if os.path.isfile(fpath):
                samples.append(
                    {
                        "filename": fname,
                        "size": os.path.getsize(fpath),
                        "ext": os.path.splitext(fname)[1].lower(),
                        "description": sample_descriptions.get(fname, "Chứng từ mẫu"),
                    }
                )

    return {
        "vessel_types": VESSEL_TYPES,
        "doc_types": DOC_TYPE_INFO,
        "sample_files": samples,
        "gdrive": get_bridge_status(),
        "gdrive_status": get_bridge_status(),
    }


@app.get("/api/vessels")
def list_vessels(current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER"))):
    vessels = get_all_vessels()
    stats = {
        "total": len(vessels),
        "red": sum(1 for v in vessels if v["status_color"] == "RED"),
        "orange": sum(1 for v in vessels if v["status_color"] == "ORANGE"),
        "green": sum(1 for v in vessels if v["status_color"] == "GREEN"),
        "total_files": sum(len(v.get("files", [])) for v in vessels),
    }
    return {"vessels": vessels, "stats": stats, "gdrive_status": get_bridge_status()}


@app.get("/api/vessels/{vessel_id}")
def get_single_vessel(
    vessel_id: int,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER")),
):
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")
    return vessel


@app.put("/api/vessels/{vessel_id}")
def edit_vessel_endpoint(
    vessel_id: int,
    payload: VesselUpdateRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    """Cập nhật Tên tàu, Loại tàu và Giờ cập cảng (ETA)."""
    existing = get_vessel_detail(vessel_id)
    if not existing:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")
    if not payload.vessel_name or not payload.vessel_name.strip():
        raise HTTPException(status_code=400, detail="Tên tàu không được để trống.")

    folder_meta = get_vessel_folder_meta_internal(vessel_id)
    updated = update_vessel_info(
        vessel_id=vessel_id,
        vessel_name=payload.vessel_name,
        eta=payload.eta,
        vessel_type=payload.vessel_type,
    )
    old_folder_name = updated.pop("_old_folder_name", folder_meta["folder_name"])
    new_folder_name = updated.pop("_new_folder_name", folder_meta["folder_name"])

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="CREATE_VESSEL",
        details=f"Cập nhật thông tin tàu: {updated['vessel_name']} (ETA: {updated['eta']})",
    )
    background_tasks.add_task(
        bg_rename_vessel_folder_and_sync,
        vessel_id,
        folder_meta["gdrive_folder_id"],
        old_folder_name,
        new_folder_name,
    )
    return {
        "message": f"Đã cập nhật thông tin tàu {updated['vessel_name']} (ETA: {updated['eta']}).",
        "vessel": updated,
    }


@app.get("/api/files/{file_id}/download")
def download_file_by_id(
    file_id: int,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER")),
):
    rec = get_file_record_internal(file_id)
    if not rec or not rec.get("abs_file_path"):
        raise HTTPException(status_code=404, detail="Không tìm thấy chứng từ.")

    abs_file = rec["abs_file_path"]
    if not abs_file.startswith(os.path.abspath(STORAGE_ROOT)):
        raise HTTPException(status_code=403, detail="Truy cập file bị từ chối.")

    if not os.path.exists(abs_file) and rec.get("gdrive_file_id"):
        download_doc_from_gdrive_if_missing(rec["gdrive_file_id"], abs_file)

    if not os.path.exists(abs_file):
        raise HTTPException(status_code=404, detail="File chứng từ không tồn tại trên bộ nhớ đệm.")

    media_type = _mime_for_ext(rec.get("file_ext") or "")
    return FileResponse(
        path=abs_file,
        filename=rec["stored_filename"],
        media_type=media_type,
    )


# ==================== EXPORT INSPECTION REPORT (PDF / DOCX / XLSX) ====================

def _status_text_vn(color: str) -> str:
    if color == "GREEN":
        return "✓ Đạt (Đủ hồ sơ 100%)"
    if color == "ORANGE":
        return "— Thiếu chứng từ (Đang bổ sung)"
    return "✕ Cần xử lý (Chưa nộp / Cảnh báo sai tàu)"


def _item_status_vn(st: str) -> str:
    if st == "COMPLETED":
        return "Đã đủ"
    if st == "INCOMPLETE":
        return "Đang thiếu"
    return "Chưa nộp"


@app.get("/api/vessels/{vessel_id}/export/pdf")
def export_vessel_checklist_pdf(
    vessel_id: int,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER")),
):
    """
    Xuất Phiếu Kiểm Tra Hồ Sơ Tàu Cập Cảng ra PDF với Font Unicode Tiếng Việt nhúng trực tiếp
    (không bao giờ lỗi dấu Tiếng Việt trên cả Windows lẫn Linux Render).
    """
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)

    font_paths = _resolve_vietnamese_fonts()
    f_reg = fitz.Font(fontfile=font_paths["regular"]) if font_paths["regular"] else fitz.Font("helv")
    f_bold = fitz.Font(fontfile=font_paths["bold"]) if font_paths["bold"] else f_reg

    def draw_icon(kind: str, x: float, y: float, color):
        """Vẽ icon ✓, —, ✕ vector sắc nét trước nhãn đối chiếu."""
        if kind == "GREEN":
            # Vector tick ✓
            page.draw_line(fitz.Point(x, y - 3.5), fitz.Point(x + 2.8, y - 0.5), color=color, width=1.4)
            page.draw_line(fitz.Point(x + 2.8, y - 0.5), fitz.Point(x + 7.5, y - 6.5), color=color, width=1.4)
        elif kind == "ORANGE":
            # Vector dash —
            page.draw_line(fitz.Point(x, y - 3.5), fitz.Point(x + 7.0, y - 3.5), color=color, width=1.5)
        else:
            # Vector cross ✕
            page.draw_line(fitz.Point(x, y - 6.5), fitz.Point(x + 6.5, y - 0.5), color=color, width=1.3)
            page.draw_line(fitz.Point(x + 6.5, y - 6.5), fitz.Point(x, y - 0.5), color=color, width=1.3)

    def draw_txt(x: float, y: float, text: str, size: float = 9.5, color=(0.06, 0.09, 0.16), bold: bool = False):
        s = str(text or "")
        offset_x = 0.0
        if s.startswith("✓ "):
            draw_icon("GREEN", x, y, color)
            s = s[2:]
            offset_x = 10.5
        elif s.startswith("— "):
            draw_icon("ORANGE", x, y, color)
            s = s[2:]
            offset_x = 10.5
        elif s.startswith("✕ "):
            draw_icon("RED", x, y, color)
            s = s[2:]
            offset_x = 10.5

        chosen_font = f_bold if bold else f_reg
        # Loại bỏ các ký tự không có glyph trong font để tránh ô vuông
        clean_chars = [ch if chosen_font.has_glyph(ord(ch)) > 0 else " " for ch in s]
        clean_s = "".join(clean_chars)

        tw = fitz.TextWriter(page.rect)
        tw.append(fitz.Point(x + offset_x, y), clean_s, font=chosen_font, fontsize=size)
        tw.write_text(page, color=color)

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    page.draw_rect(fitz.Rect(36, 32, 559, 92), color=(0.06, 0.09, 0.16), fill=(0.97, 0.98, 0.99), width=1)
    draw_txt(48, 56, "PHIẾU KIỂM TRA HỒ SƠ TÀU CẬP CẢNG", size=13.5, color=(0.06, 0.09, 0.16), bold=True)
    draw_txt(48, 76, "PORT-CALL DOCUMENT VERIFICATION REPORT · PORTCALL DOCS", size=8.5, color=(0.39, 0.45, 0.55))

    y = 115
    draw_txt(40, y, f"Tên tàu: {vessel['vessel_name']}", size=10.5, bold=True)
    draw_txt(310, y, f"Loại tàu: {vessel['vessel_type_label']}", size=10)
    y += 20
    draw_txt(40, y, f"Giờ dự kiến đến (ETA): {vessel['eta']}", size=9.5)
    draw_txt(310, y, f"Tổng số B/L theo Manifest: {vessel.get('expected_bl_count', 0)} B/L", size=9.5)
    y += 20

    st_color = (0.06, 0.72, 0.50) if vessel["status_color"] == "GREEN" else (
        (0.92, 0.55, 0.03) if vessel["status_color"] == "ORANGE" else (0.90, 0.22, 0.22)
    )
    draw_txt(40, y, f"Trạng thái tổng thể: {_status_text_vn(vessel['status_color'])}", size=10, color=st_color, bold=True)
    draw_txt(310, y, f"Tiến độ: {vessel['completed_checklist_count']}/{vessel['total_checklist_count']} hạng mục", size=9.5)
    y += 20
    draw_txt(
        40,
        y,
        f"Người kiểm tra: {current_user['full_name']} ({current_user['username']} · {current_user['role']})",
        size=9,
        color=(0.28, 0.33, 0.41),
    )
    draw_txt(380, y, f"Ngày xuất: {now_str}", size=9, color=(0.28, 0.33, 0.41))

    y += 24
    page.draw_rect(fitz.Rect(36, y, 559, y + 24), color=(0.82, 0.85, 0.90), fill=(0.94, 0.96, 0.98), width=0.8)
    draw_txt(42, y + 16, "MÃ CHUẨN", size=8.5, color=(0.15, 0.20, 0.28), bold=True)
    draw_txt(136, y + 16, "LOẠI CHỨNG TỪ", size=8.5, color=(0.15, 0.20, 0.28), bold=True)
    draw_txt(265, y + 16, "ĐỐI CHIẾU", size=8.5, color=(0.15, 0.20, 0.28), bold=True)
    draw_txt(385, y + 16, "CHỨNG TỪ ĐÃ NỘP", size=8.5, color=(0.15, 0.20, 0.28), bold=True)
    y += 24

    for item in vessel.get("checklist", []):
        files = item.get("files", [])
        row_h = max(28, 17 * max(1, len(files)) + 8)
        page.draw_rect(fitz.Rect(36, y, 559, y + row_h), color=(0.88, 0.91, 0.94), fill=None, width=0.6)

        m_state = item.get("match_state", "RED")
        m_label = item.get("match_label", "✕ Chưa nộp")
        c_st = (0.06, 0.72, 0.50) if m_state == "GREEN" else (
            (0.92, 0.55, 0.03) if m_state == "ORANGE" else (0.90, 0.22, 0.22)
        )

        draw_txt(42, y + 17, item["doc_type"], size=8.5, bold=True)
        draw_txt(136, y + 17, item["doc_name"][:25], size=9)
        draw_txt(265, y + 17, m_label, size=9, color=c_st, bold=True)

        if not files:
            draw_txt(385, y + 17, "Chưa nộp", size=8.5, color=(0.55, 0.6, 0.68))
        else:
            fy = y + 16
            for f in files:
                draw_txt(385, fy, f"{f['stored_filename'][:29]} ({f.get('uploaded_by', 'ops')})", size=8)
                fy += 15

        y += row_h

    wrong_files = vessel.get("wrong_vessel_files", [])
    unclass_files = vessel.get("unclassified_files", [])
    if wrong_files or unclass_files:
        y += 18
        draw_txt(40, y, "CẢNH BÁO KIỂM TRA CHÉO:", size=9.5, color=(0.90, 0.22, 0.22), bold=True)
        y += 16
        for wf in wrong_files:
            draw_txt(46, y, f"✕ [SAI TÀU] {wf['original_filename']}: {wf['error_message']}", size=8.5, color=(0.90, 0.22, 0.22))
            y += 15
        for uf in unclass_files:
            draw_txt(46, y, f"— [CHƯA PHÂN LOẠI] {uf['original_filename']}: Chờ chọn loại chứng từ", size=8.5, color=(0.92, 0.55, 0.03))
            y += 15

    sig_y = min(740, max(y + 45, 650))
    draw_txt(65, sig_y, "ĐẠI DIỆN ĐẠI LÝ HÃNG TÀU", size=9.5, bold=True)
    draw_txt(65, sig_y + 15, "(Ký, ghi rõ họ tên)", size=8.5, color=(0.45, 0.5, 0.58))

    draw_txt(360, sig_y, "CHUYÊN VIÊN KIỂM TRA HỒ SƠ", size=9.5, bold=True)
    draw_txt(360, sig_y + 15, "(Ký, ghi rõ họ tên)", size=8.5, color=(0.45, 0.5, 0.58))
    draw_txt(360, sig_y + 62, current_user["full_name"], size=9.5, bold=True)

    pdf_bytes = doc.tobytes(garbage=3, deflate=True)
    doc.close()

    safe_vname = sanitize_token(vessel["vessel_name"])
    filename = f"Phieu_Kiem_Tra_{safe_vname}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/vessels/{vessel_id}/export/docx")
def export_vessel_checklist_docx(
    vessel_id: int,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER")),
):
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    doc = docx.Document()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("PHIẾU KIỂM TRA HỒ SƠ TÀU CẬP CẢNG\n")
    r_title.bold = True
    r_title.font.size = Pt(15)
    r_sub = title_p.add_run("PORT-CALL DOCUMENT VERIFICATION REPORT")
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    info_p = doc.add_paragraph()
    info_p.add_run("Tên tàu: ").bold = True
    info_p.add_run(f"{vessel['vessel_name']}\t\t")
    info_p.add_run("Loại tàu: ").bold = True
    info_p.add_run(f"{vessel['vessel_type_label']}\n")

    info_p.add_run("Giờ dự kiến cập cảng (ETA): ").bold = True
    info_p.add_run(f"{vessel['eta']}\t\t")
    info_p.add_run("Tổng B/L theo Manifest: ").bold = True
    info_p.add_run(f"{vessel.get('expected_bl_count', 0)} B/L\n")

    info_p.add_run("Trạng thái tổng thể: ").bold = True
    info_p.add_run(f"{_status_text_vn(vessel['status_color'])} ({vessel['completed_checklist_count']}/{vessel['total_checklist_count']} hạng mục)\n")

    info_p.add_run("Người kiểm tra: ").bold = True
    info_p.add_run(f"{current_user['full_name']} ({current_user['username']} - {current_user['role']})\t\t")
    info_p.add_run("Thời gian xuất: ").bold = True
    info_p.add_run(now_str)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers_text = ["Mã chuẩn", "Loại chứng từ", "Đối chiếu", "Chứng từ đã nộp & Người tải lên"]
    for idx, h_txt in enumerate(headers_text):
        hdr[idx].text = h_txt
        for run in hdr[idx].paragraphs[0].runs:
            run.bold = True

    for item in vessel.get("checklist", []):
        row_cells = table.add_row().cells
        row_cells[0].text = item["doc_type"]
        row_cells[1].text = item["doc_name"]

        m_label = item.get("match_label", "✕ Chưa nộp")
        m_state = item.get("match_state", "RED")
        r_m = row_cells[2].paragraphs[0].add_run(m_label)
        r_m.bold = True
        if m_state == "GREEN":
            r_m.font.color.rgb = RGBColor(16, 185, 129)
        elif m_state == "ORANGE":
            r_m.font.color.rgb = RGBColor(245, 158, 11)
        else:
            r_m.font.color.rgb = RGBColor(239, 68, 68)

        files = item.get("files", [])
        if not files:
            row_cells[3].text = "— Chưa nộp"
        else:
            file_lines = [
                f"{f['stored_filename']} (gốc: {f['original_filename']} · {f.get('uploaded_by', 'ops')} · {f.get('uploaded_at', '')})"
                for f in files
            ]
            row_cells[3].text = "\n".join(file_lines)

    wrong_files = vessel.get("wrong_vessel_files", [])
    unclass_files = vessel.get("unclassified_files", [])
    if wrong_files or unclass_files:
        wp = doc.add_paragraph()
        w_run = wp.add_run("\nCẢNH BÁO LỖI HỒ SƠ:\n")
        w_run.bold = True
        w_run.font.color.rgb = RGBColor(220, 38, 38)
        for wf in wrong_files:
            wp.add_run(f"• [SAI TÀU] {wf['original_filename']}: {wf['error_message']}\n")
        for uf in unclass_files:
            wp.add_run(f"• [CHƯA PHÂN LOẠI] {uf['original_filename']}: Chưa nhận diện được loại chứng từ.\n")

    sig_table = doc.add_table(rows=1, cols=2)
    s_cells = sig_table.rows[0].cells
    p_left = s_cells[0].paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.add_run("\nĐẠI DIỆN ĐẠI LÝ HÃNG TÀU\n").bold = True
    p_left.add_run("(Ký, ghi rõ họ tên)\n\n\n")

    p_right = s_cells[1].paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.add_run("\nCHUYÊN VIÊN KIỂM TRA HỒ SƠ\n").bold = True
    p_right.add_run("(Ký, ghi rõ họ tên)\n\n\n")
    p_right.add_run(current_user["full_name"]).bold = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    safe_vname = sanitize_token(vessel["vessel_name"])
    filename = f"Phieu_Kiem_Tra_{safe_vname}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/vessels/{vessel_id}/export/xlsx")
def export_vessel_checklist_xlsx(
    vessel_id: int,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER")),
):
    """Xuất Phiếu Kiểm Tra Hồ Sơ Tàu Cập Cảng ra bảng Excel (.xlsx) bằng openpyxl."""
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Phieu_Kiem_Tra"

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.merge_cells("A1:E1")
    ws["A1"] = "PHIẾU KIỂM TRA HỒ SƠ TÀU CẬP CẢNG (PORTCALL DOCS)"
    ws["A1"].font = XlFont(name="Arial", size=14, bold=True, color="0F172A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"] = "Tên tàu:"
    ws["A3"].font = XlFont(bold=True)
    ws["B3"] = vessel["vessel_name"]
    ws["C3"] = "Loại tàu:"
    ws["C3"].font = XlFont(bold=True)
    ws["D3"] = vessel["vessel_type_label"]

    ws["A4"] = "Giờ cập cảng (ETA):"
    ws["A4"].font = XlFont(bold=True)
    ws["B4"] = vessel["eta"]
    ws["C4"] = "Tổng B/L theo Manifest:"
    ws["C4"].font = XlFont(bold=True)
    ws["D4"] = f"{vessel.get('expected_bl_count', 0)} B/L"

    ws["A5"] = "Trạng thái tổng thể:"
    ws["A5"].font = XlFont(bold=True)
    ws["B5"] = _status_text_vn(vessel["status_color"])
    ws["C5"] = "Người kiểm tra:"
    ws["C5"].font = XlFont(bold=True)
    ws["D5"] = f"{current_user['full_name']} ({current_user['username']}) - {now_str}"

    headers = ["Mã chuẩn", "Loại chứng từ", "Đối chiếu", "Chứng từ đã nộp", "Người tải lên & Thời gian"]
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx, value=h_text)
        cell.font = XlFont(name="Arial", size=10, bold=True, color="0F172A")
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    row_num = 8
    for item in vessel.get("checklist", []):
        files = item.get("files", [])
        m_state = item.get("match_state", "RED")
        m_label = item.get("match_label", "✕ Chưa nộp")

        c1 = ws.cell(row=row_num, column=1, value=item["doc_type"])
        c2 = ws.cell(row=row_num, column=2, value=item["doc_name"])
        c3 = ws.cell(row=row_num, column=3, value=m_label)

        if files:
            c4 = ws.cell(row=row_num, column=4, value="\n".join(f["stored_filename"] for f in files))
            c5 = ws.cell(
                row=row_num,
                column=5,
                value="\n".join(f"{f.get('uploaded_by', 'ops')} ({f.get('uploaded_at', '')})" for f in files),
            )
        else:
            c4 = ws.cell(row=row_num, column=4, value="— Chưa nộp")
            c5 = ws.cell(row=row_num, column=5, value="—")

        color_hex = "10B981" if m_state == "GREEN" else ("F59E0B" if m_state == "ORANGE" else "EF4444")
        c1.font = XlFont(name="Arial", size=10, bold=True)
        c3.font = XlFont(name="Arial", size=10, bold=True, color=color_hex)

        for c in (c1, c2, c3, c4, c5):
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)

        row_num += 1

    col_widths = [18, 28, 24, 38, 30]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_vname = sanitize_token(vessel["vessel_name"])
    filename = f"Phieu_Kiem_Tra_{safe_vname}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ==================== VESSEL MUTATION ENDPOINTS & 2-WAY GDRIVE SYNC ====================

@app.post("/api/vessels/batch")
def create_vessels_batch(
    payload: VesselBatchCreateRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    if not payload.vessels:
        raise HTTPException(status_code=400, detail="Danh sách tàu cần tạo không được để trống.")

    created = []
    for item in payload.vessels:
        if not item.vessel_name or not item.vessel_name.strip():
            continue
        clean_eta = (item.eta or datetime.now().strftime("%Y-%m-%d %H:%M")).replace("T", " ")
        v = create_vessel(
            vessel_name=item.vessel_name,
            eta=clean_eta,
            vessel_type=item.vessel_type or "CONTAINER_INTL",
        )
        created.append(v)
        folder_meta = get_vessel_folder_meta_internal(v["id"])
        background_tasks.add_task(bg_create_vessel_folder_and_sync, v["id"], folder_meta["folder_name"])

    if not created:
        raise HTTPException(status_code=400, detail="Vui lòng nhập tên tàu hợp lệ.")

    names_str = ", ".join([v["vessel_name"] for v in created])
    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="CREATE_VESSEL",
        details=f"Tạo {len(created)} lịch tàu: {names_str}",
    )

    return {
        "message": f"Đã tạo {len(created)} lịch tàu.",
        "created_vessels": created,
    }


@app.delete("/api/vessels/{vessel_id}")
def delete_vessel(
    vessel_id: int,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN")),
):
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    folder_meta = get_vessel_folder_meta_internal(vessel_id)
    abs_folder = get_vessel_storage_abs_path(vessel_id)
    if abs_folder and os.path.exists(abs_folder):
        shutil.rmtree(abs_folder, ignore_errors=True)

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM uploaded_files WHERE vessel_id = ?", (vessel_id,))
    cursor.execute("DELETE FROM checklist_items WHERE vessel_id = ?", (vessel_id,))
    cursor.execute("DELETE FROM vessels WHERE id = ?", (vessel_id,))
    conn.commit()
    conn.close()

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="DELETE_FILE",
        details=f"Xóa tàu {vessel['vessel_name']} và toàn bộ chứng từ liên quan",
    )
    background_tasks.add_task(
        bg_delete_vessel_folder_and_sync,
        folder_meta["gdrive_folder_id"],
        folder_meta["folder_name"],
    )

    return {"message": f"Đã xóa tàu {vessel['vessel_name']}."}


@app.delete("/api/vessels/{vessel_id}/clear-files")
def clear_vessel_files(
    vessel_id: int,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    folder_meta = get_vessel_folder_meta_internal(vessel_id)
    abs_folder = get_vessel_storage_abs_path(vessel_id)
    if abs_folder and os.path.exists(abs_folder):
        for fname in os.listdir(abs_folder):
            fpath = os.path.join(abs_folder, fname)
            if os.path.isfile(fpath):
                try:
                    os.remove(fpath)
                except Exception:
                    pass

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM uploaded_files WHERE vessel_id = ?", (vessel_id,))
    cursor.execute("UPDATE checklist_items SET is_manual_verified = 0 WHERE vessel_id = ?", (vessel_id,))
    conn.commit()
    conn.close()

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="DELETE_FILE",
        details=f"Làm trống toàn bộ hồ sơ của tàu {vessel['vessel_name']}",
    )

    updated_vessel = recalculate_vessel_status(vessel_id)
    background_tasks.add_task(
        bg_clear_vessel_folder_and_sync,
        folder_meta["gdrive_folder_id"],
        folder_meta["folder_name"],
    )

    return {
        "message": f"Đã làm trống hồ sơ tàu {vessel['vessel_name']}.",
        "vessel": updated_vessel,
    }


@app.post("/api/vessels/{vessel_id}/checklist/{doc_type}/verify")
def verify_checklist_item_manual(
    vessel_id: int,
    doc_type: str,
    payload: ChecklistVerifyRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    """Cho phép User chủ động bấm '✓ Xác nhận Đạt' (Manual Override) trên dòng Checklist."""
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    target_type = doc_type.strip().upper()
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE checklist_items SET is_manual_verified = ? WHERE vessel_id = ? AND doc_type = ?",
        (1 if payload.verified else 0, vessel_id, target_type),
    )
    conn.commit()
    conn.close()

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="UPLOAD_FILE",
        details=f"Tàu {vessel['vessel_name']}: {'Xác nhận Đạt (✓)' if payload.verified else 'Hủy xác nhận thủ công'} hạng mục {target_type}",
    )

    updated_vessel = recalculate_vessel_status(vessel_id)
    background_tasks.add_task(sync_db_to_gdrive)

    return {
        "message": f"Đã {'xác nhận ✓ Đạt' if payload.verified else 'chuyển về đối chiếu tự động'} cho hạng mục {target_type}.",
        "vessel": updated_vessel,
    }


@app.post("/api/vessels/{vessel_id}/upload")
async def upload_files_to_vessel(
    vessel_id: int,
    background_tasks: BackgroundTasks,
    files: List[UploadFile] = File(...),
    replace_existing: bool = Form(False),
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    abs_vessel_folder = get_vessel_storage_abs_path(vessel_id)
    if not abs_vessel_folder:
        raise HTTPException(status_code=500, detail="Không thể truy cập vùng lưu trữ của tàu.")

    folder_meta = get_vessel_folder_meta_internal(vessel_id)
    processed_results = []
    conflicts = []

    for upload_file in files:
        content = await upload_file.read()
        clean_orig_name = validate_file_security(upload_file.filename or "document.pdf", content)

        temp_name = f"__tmp_{datetime.now().strftime('%H%M%S%f')}_{clean_orig_name}"
        temp_path = os.path.join(abs_vessel_folder, temp_name)

        with open(temp_path, "wb") as f_out:
            f_out.write(content)

        res = _process_and_store_single_file(
            vessel=vessel,
            temp_file_path=temp_path,
            original_filename=clean_orig_name,
            uploaded_by=current_user["username"],
            replace_existing=replace_existing,
        )
        processed_results.append(res)

        if res.get("replaced_gdrive_items"):
            background_tasks.add_task(bg_batch_delete_docs_and_sync, res["replaced_gdrive_items"])

        if res.get("has_same_ext_conflict"):
            conflicts.append(
                {
                    "file_id": res["id"],
                    "stored_filename": res["stored_filename"],
                    "base_standard_filename": res["base_standard_filename"],
                    "doc_type": res["doc_type"],
                    "ext": res["ext"],
                }
            )

        log_audit_event(
            username=current_user["username"],
            role=current_user["role"],
            action="UPLOAD_FILE",
            details=(
                f"Tàu {vessel['vessel_name']}: Tải lên '{clean_orig_name}' -> "
                f"'{res['stored_filename']}' ({res['doc_type']})"
                + (" [TẠM GIỮ CHỜ XÁC NHẬN - KHÔNG LƯU DRIVE]" if res["is_warning"] else "")
            ),
        )

        # CHẶN LƯU FILE LỖI LÊN DRIVE: Chỉ đẩy lên Drive khi file KHÔNG bị cảnh báo WRONG_VESSEL / UNCLASSIFIED
        if not res["is_warning"]:
            background_tasks.add_task(
                bg_upload_doc_and_sync,
                res["id"],
                vessel_id,
                folder_meta["folder_name"],
                res["abs_path"],
                res["stored_filename"],
                _mime_for_ext(res["ext"]),
            )
        else:
            background_tasks.add_task(sync_db_to_gdrive)

    updated_vessel = recalculate_vessel_status(vessel_id)
    return {
        "message": f"Đã xử lý {len(processed_results)} chứng từ.",
        "processed_files": processed_results,
        "conflicts": conflicts,
        "vessel": updated_vessel,
    }


@app.post("/api/vessels/{vessel_id}/simulate-sample")
def simulate_sample_upload(
    vessel_id: int,
    payload: SimulateSampleRequest,
    background_tasks: BackgroundTasks = BackgroundTasks(),
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    abs_vessel_folder = get_vessel_storage_abs_path(vessel_id)
    if not abs_vessel_folder:
        raise HTTPException(status_code=500, detail="Không thể truy cập vùng lưu trữ của tàu.")

    folder_meta = get_vessel_folder_meta_internal(vessel_id)
    processed_results = []
    conflicts = []

    for fname in payload.filenames:
        if fname == "__UNCLASSIFIED_DEMO__":
            orig_name = "scan_anh_mo_009.pdf"
            temp_path = os.path.join(abs_vessel_folder, f"__tmp_{datetime.now().strftime('%H%M%S%f')}_{orig_name}")
            create_pdf_file(
                temp_path,
                title="SCANNED DOCUMENT - PAGE 01",
                subtitle="Optical Scan Image - Low Contrast Maritime Attachment",
                lines=[
                    f"Vessel Name: {vessel['vessel_name']}",
                    "Reference ID: DOC-2026-99812",
                    "===",
                    "This scanned page contains no standard classification anchor keywords.",
                ],
            )
            res = _process_and_store_single_file(
                vessel=vessel,
                temp_file_path=temp_path,
                original_filename=orig_name,
                uploaded_by=current_user["username"],
                replace_existing=bool(payload.replace_existing),
            )
            processed_results.append(res)
            log_audit_event(
                username=current_user["username"],
                role=current_user["role"],
                action="UPLOAD_FILE",
                details=f"Tàu {vessel['vessel_name']}: Nạp file mẫu '{orig_name}' (Chờ xác nhận, chưa lưu Drive)",
            )
            # Không lưu file UNCLASSIFIED lên Drive cho tới khi User xác nhận
            background_tasks.add_task(sync_db_to_gdrive)
            continue

        clean_fname = sanitize_filename(fname)
        sample_src = os.path.join(SAMPLE_DIR, clean_fname)
        if not os.path.exists(sample_src):
            raise HTTPException(status_code=404, detail=f"Không tìm thấy chứng từ mẫu '{clean_fname}'.")

        with open(sample_src, "rb") as f_in:
            content = f_in.read()
        validate_file_security(clean_fname, content)

        temp_path = os.path.join(abs_vessel_folder, f"__tmp_{datetime.now().strftime('%H%M%S%f')}_{clean_fname}")
        with open(temp_path, "wb") as f_out:
            f_out.write(content)

        res = _process_and_store_single_file(
            vessel=vessel,
            temp_file_path=temp_path,
            original_filename=clean_fname,
            uploaded_by=current_user["username"],
            replace_existing=bool(payload.replace_existing),
        )
        processed_results.append(res)

        if res.get("has_same_ext_conflict"):
            conflicts.append(
                {
                    "file_id": res["id"],
                    "stored_filename": res["stored_filename"],
                    "base_standard_filename": res["base_standard_filename"],
                    "doc_type": res["doc_type"],
                    "ext": res["ext"],
                }
            )

        log_audit_event(
            username=current_user["username"],
            role=current_user["role"],
            action="UPLOAD_FILE",
            details=(
                f"Tàu {vessel['vessel_name']}: Nạp chứng từ '{clean_fname}' -> "
                f"'{res['stored_filename']}' ({res['doc_type']})"
                + (" [TẠM GIỮ CHỜ XÁC NHẬN - KHÔNG LƯU DRIVE]" if res["is_warning"] else "")
            ),
        )
        # Chỉ đẩy lên Google Drive khi file hợp lệ (không bị WRONG_VESSEL / UNCLASSIFIED)
        if not res["is_warning"]:
            background_tasks.add_task(
                bg_upload_doc_and_sync,
                res["id"],
                vessel_id,
                folder_meta["folder_name"],
                res["abs_path"],
                res["stored_filename"],
                _mime_for_ext(res["ext"]),
            )
        else:
            background_tasks.add_task(sync_db_to_gdrive)

    updated_vessel = recalculate_vessel_status(vessel_id)
    return {
        "message": f"Đã nạp {len(processed_results)} chứng từ mẫu.",
        "processed_files": processed_results,
        "conflicts": conflicts,
        "vessel": updated_vessel,
    }


@app.put("/api/files/{file_id}/classify")
def manual_classify_and_confirm_file(
    file_id: int,
    payload: ClassifyFileRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    """
    Xác nhận chứng từ là đúng (Override lỗi WRONG_VESSEL hoặc phân loại UNCLASSIFIED),
    đổi tên chuẩn [{LOAI_GIAY_TO}]_{Ten_Tau}.{ext}, hỗ trợ ghi đè file cũ cùng định dạng,
    và đẩy lên Google Drive!
    """
    new_doc_type = payload.doc_type.strip().upper()
    if new_doc_type not in DOC_TYPE_INFO or new_doc_type == "UNCLASSIFIED":
        raise HTTPException(status_code=400, detail="Loại chứng từ chọn không hợp lệ.")

    rec = get_file_record_internal(file_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Không tìm thấy chứng từ.")

    vessel_id = rec["vessel_id"]
    abs_vessel_folder = rec["abs_vessel_folder"]
    old_abs_path = rec["abs_file_path"]
    ext = (rec["file_ext"] or os.path.splitext(rec["original_filename"])[1] or ".pdf").lower()
    folder_meta = get_vessel_folder_meta_internal(vessel_id)

    conn = get_connection()
    cursor = conn.cursor()

    # Kiểm tra các file hợp lệ khác cùng (vessel_id, new_doc_type, ext)
    cursor.execute(
        """
        SELECT id, stored_filename, gdrive_file_id
        FROM uploaded_files
        WHERE vessel_id = ? AND doc_type = ? AND LOWER(file_ext) = ? AND is_wrong_vessel = 0 AND id != ?
        ORDER BY id ASC
        """,
        (vessel_id, new_doc_type, ext, file_id),
    )
    existing_same_ext = [dict(r) for r in cursor.fetchall()]

    deleted_drive_items: List[Dict[str, str]] = []
    if existing_same_ext and payload.replace_existing:
        for old_f in existing_same_ext:
            old_p = os.path.join(abs_vessel_folder, os.path.basename(old_f["stored_filename"]))
            if os.path.exists(old_p) and os.path.abspath(old_p) != os.path.abspath(old_abs_path):
                try:
                    os.remove(old_p)
                except Exception:
                    pass
            cursor.execute("DELETE FROM uploaded_files WHERE id = ?", (old_f["id"],))
            deleted_drive_items.append(
                {
                    "gdrive_file_id": old_f.get("gdrive_file_id") or "",
                    "vessel_folder_name": folder_meta["folder_name"],
                    "stored_filename": old_f["stored_filename"],
                }
            )
        seq_num = 1
    else:
        seq_num = len(existing_same_ext) + 1

    new_stored_filename = build_standardized_filename(new_doc_type, rec["vessel_name"], seq_num, ext)
    new_abs_path = os.path.join(abs_vessel_folder, new_stored_filename)

    while os.path.exists(new_abs_path) and os.path.abspath(new_abs_path) != os.path.abspath(old_abs_path):
        seq_num += 1
        new_stored_filename = build_standardized_filename(new_doc_type, rec["vessel_name"], seq_num, ext)
        new_abs_path = os.path.join(abs_vessel_folder, new_stored_filename)

    if os.path.exists(old_abs_path) and os.path.abspath(old_abs_path) != os.path.abspath(new_abs_path):
        shutil.move(old_abs_path, new_abs_path)

    target_parse_path = new_abs_path if os.path.exists(new_abs_path) else old_abs_path
    parsed = parse_and_validate_file(
        file_path=target_parse_path,
        original_filename=rec["original_filename"],
        current_vessel_name=rec["vessel_name"],
        forced_doc_type=new_doc_type,
    )

    # Nếu User bấm "Vẫn xác nhận là đúng & Lưu Drive" (confirm_override=True), xóa bỏ cờ lỗi sai tàu
    final_is_wrong = 0 if payload.confirm_override else parsed["is_wrong_vessel"]
    final_err_code = "" if payload.confirm_override else parsed["error_code"]
    final_err_msg = "" if payload.confirm_override else parsed["error_message"]

    cursor.execute(
        """
        UPDATE uploaded_files
        SET stored_filename = ?,
            file_path = ?,
            doc_type = ?,
            detected_method = 'MANUAL_CONFIRMED',
            detected_keyword = ?,
            detected_vessel_name = ?,
            is_wrong_vessel = ?,
            error_code = ?,
            error_message = ?,
            extracted_count = ?,
            extracted_codes_json = ?,
            sync_status = ?
        WHERE id = ?
        """,
        (
            new_stored_filename,
            new_stored_filename,
            new_doc_type,
            f"User Confirmed: {new_doc_type}",
            parsed["detected_vessel_name"],
            final_is_wrong,
            final_err_code,
            final_err_msg,
            parsed["extracted_count"],
            json.dumps(parsed["extracted_codes"], ensure_ascii=False),
            "PENDING_SYNC" if is_bridge_enabled() else "LOCAL",
            file_id,
        ),
    )

    if payload.mark_checklist_completed:
        cursor.execute(
            "UPDATE checklist_items SET is_manual_verified = 1 WHERE vessel_id = ? AND doc_type = ?",
            (vessel_id, new_doc_type),
        )

    conn.commit()
    conn.close()

    if deleted_drive_items:
        background_tasks.add_task(bg_batch_delete_docs_and_sync, deleted_drive_items)

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="UPLOAD_FILE",
        details=f"Xác nhận & Lưu Drive chứng từ '{rec['original_filename']}' -> '{new_stored_filename}' ({new_doc_type})",
    )

    updated_vessel = recalculate_vessel_status(vessel_id)
    background_tasks.add_task(
        bg_rename_doc_and_sync,
        file_id,
        vessel_id,
        rec.get("gdrive_file_id", ""),
        new_stored_filename,
        folder_meta["folder_name"],
        new_abs_path,
        _mime_for_ext(ext),
        bool(payload.replace_existing),
    )

    return {
        "message": f"Đã xác nhận '{new_stored_filename}' và đồng bộ lên Google Drive.",
        "vessel": updated_vessel,
    }


@app.post("/api/files/{file_id}/replace-existing")
def replace_existing_same_ext_file(
    file_id: int,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    """
    Khi người dùng tải lên file thứ 2 cùng loại chứng từ và cùng định dạng (đang mang hậu tố _(2)),
    cho phép bấm 1-Click 'Ghi đè (Thay thế file cũ trên Drive)' để xóa file cũ và đổi tên về chuẩn gốc.
    """
    rec = get_file_record_internal(file_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Không tìm thấy chứng từ.")

    vessel_id = rec["vessel_id"]
    doc_type = rec["doc_type"]
    ext = (rec["file_ext"] or ".pdf").lower()
    abs_vessel_folder = rec["abs_vessel_folder"]
    current_abs_path = rec["abs_file_path"]
    folder_meta = get_vessel_folder_meta_internal(vessel_id)

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT id, stored_filename, gdrive_file_id
        FROM uploaded_files
        WHERE vessel_id = ? AND doc_type = ? AND LOWER(file_ext) = ? AND id != ?
        """,
        (vessel_id, doc_type, ext, file_id),
    )
    older_files = [dict(r) for r in cursor.fetchall()]

    deleted_drive_items: List[Dict[str, str]] = []
    for old_f in older_files:
        old_p = os.path.join(abs_vessel_folder, os.path.basename(old_f["stored_filename"]))
        if os.path.exists(old_p) and os.path.abspath(old_p) != os.path.abspath(current_abs_path):
            try:
                os.remove(old_p)
            except Exception:
                pass
        cursor.execute("DELETE FROM uploaded_files WHERE id = ?", (old_f["id"],))
        deleted_drive_items.append(
            {
                "gdrive_file_id": old_f.get("gdrive_file_id") or "",
                "vessel_folder_name": folder_meta["folder_name"],
                "stored_filename": old_f["stored_filename"],
            }
        )

    base_filename = build_standardized_filename(doc_type, rec["vessel_name"], 1, ext)
    new_abs_path = os.path.join(abs_vessel_folder, base_filename)
    if os.path.exists(current_abs_path) and os.path.abspath(current_abs_path) != os.path.abspath(new_abs_path):
        shutil.move(current_abs_path, new_abs_path)

    cursor.execute(
        "UPDATE uploaded_files SET stored_filename = ?, file_path = ? WHERE id = ?",
        (base_filename, base_filename, file_id),
    )
    conn.commit()
    conn.close()

    if deleted_drive_items:
        background_tasks.add_task(bg_batch_delete_docs_and_sync, deleted_drive_items)

    background_tasks.add_task(
        bg_rename_doc_and_sync,
        file_id,
        vessel_id,
        rec.get("gdrive_file_id", ""),
        base_filename,
        folder_meta["folder_name"],
        new_abs_path,
        _mime_for_ext(ext),
        True,
    )

    updated_vessel = recalculate_vessel_status(vessel_id)
    return {
        "message": f"Đã ghi đè thành công bằng '{base_filename}' trên Web và Google Drive.",
        "vessel": updated_vessel,
    }


@app.delete("/api/files/{file_id}")
def delete_uploaded_file(
    file_id: int,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    rec = get_file_record_internal(file_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Không tìm thấy chứng từ cần xóa.")

    vessel_id = rec["vessel_id"]
    doc_type = rec["doc_type"]
    abs_file = rec.get("abs_file_path")
    if abs_file and os.path.exists(abs_file):
        try:
            os.remove(abs_file)
        except Exception:
            pass

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM uploaded_files WHERE id = ?", (file_id,))
    cursor.execute(
        "SELECT COUNT(*) AS cnt FROM uploaded_files WHERE vessel_id = ? AND doc_type = ? AND is_wrong_vessel = 0",
        (vessel_id, doc_type),
    )
    rem = cursor.fetchone()
    if not rem or rem["cnt"] == 0:
        cursor.execute(
            "UPDATE checklist_items SET is_manual_verified = 0 WHERE vessel_id = ? AND doc_type = ?",
            (vessel_id, doc_type),
        )
    conn.commit()
    conn.close()

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="DELETE_FILE",
        details=f"Xóa chứng từ '{rec['stored_filename']}' (gốc: '{rec['original_filename']}') khỏi tàu {rec.get('vessel_name', '')}",
    )

    updated_vessel = recalculate_vessel_status(vessel_id)
    background_tasks.add_task(
        bg_delete_doc_and_sync,
        rec.get("gdrive_file_id", ""),
        rec.get("vessel_folder_name", ""),
        rec["stored_filename"],
    )

    return {
        "message": f"Đã gỡ chứng từ '{rec['stored_filename']}'.",
        "vessel": updated_vessel,
    }


# ==================== BATCH FILE ACTIONS (SELECT ALL / MULTI-SELECT) ====================

@app.post("/api/files/batch-delete")
def batch_delete_files(
    payload: BatchDeleteFilesRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    """Gỡ hàng loạt các file đã tích chọn trên cả Web và Google Drive."""
    if not payload.file_ids:
        raise HTTPException(status_code=400, detail="Chưa chọn file nào để gỡ.")

    affected_vessel_ids = set()
    drive_delete_list: List[Dict[str, str]] = []
    deleted_names = []

    conn = get_connection()
    cursor = conn.cursor()

    for fid in payload.file_ids:
        rec = get_file_record_internal(fid)
        if not rec:
            continue
        affected_vessel_ids.add(rec["vessel_id"])
        abs_file = rec.get("abs_file_path")
        if abs_file and os.path.exists(abs_file):
            try:
                os.remove(abs_file)
            except Exception:
                pass
        cursor.execute("DELETE FROM uploaded_files WHERE id = ?", (fid,))
        deleted_names.append(rec["stored_filename"])
        drive_delete_list.append(
            {
                "gdrive_file_id": rec.get("gdrive_file_id") or "",
                "vessel_folder_name": rec.get("vessel_folder_name") or "",
                "stored_filename": rec["stored_filename"],
            }
        )

    conn.commit()
    conn.close()

    updated_vessel = None
    for vid in affected_vessel_ids:
        updated_vessel = recalculate_vessel_status(vid)

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="DELETE_FILE",
        details=f"Gỡ hàng loạt {len(deleted_names)} chứng từ: {', '.join(deleted_names[:5])}",
    )
    background_tasks.add_task(bg_batch_delete_docs_and_sync, drive_delete_list)

    return {
        "message": f"Đã gỡ {len(deleted_names)} chứng từ đã chọn trên Web và Google Drive.",
        "vessel": updated_vessel,
    }


@app.post("/api/files/batch-confirm")
def batch_confirm_files(
    payload: BatchConfirmFilesRequest,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    """Xác nhận Đạt (✓) hàng loạt các file đã tích chọn và đẩy các file đang chờ lên Google Drive."""
    if not payload.file_ids:
        raise HTTPException(status_code=400, detail="Chưa chọn file nào để xác nhận.")

    affected_vessel_ids = set()
    confirmed_count = 0

    for fid in payload.file_ids:
        rec = get_file_record_internal(fid)
        if not rec:
            continue
        vid = rec["vessel_id"]
        affected_vessel_ids.add(vid)
        target_doc_type = (payload.doc_type or rec["doc_type"] or "MANIFEST").strip().upper()
        if target_doc_type == "UNCLASSIFIED":
            target_doc_type = "MANIFEST"

        abs_vessel_folder = rec["abs_vessel_folder"]
        old_abs_path = rec["abs_file_path"]
        ext = (rec["file_ext"] or ".pdf").lower()
        folder_meta = get_vessel_folder_meta_internal(vid)

        conn = get_connection()
        cursor = conn.cursor()

        # Nếu file đang là UNCLASSIFIED hoặc WRONG_VESSEL thì đổi tên chuẩn và gỡ lỗi
        if rec["is_wrong_vessel"] == 1 or rec["doc_type"] == "UNCLASSIFIED" or (payload.doc_type and payload.doc_type != rec["doc_type"]):
            cursor.execute(
                """
                SELECT COUNT(*) AS cnt FROM uploaded_files
                WHERE vessel_id = ? AND doc_type = ? AND LOWER(file_ext) = ? AND is_wrong_vessel = 0 AND id != ?
                """,
                (vid, target_doc_type, ext, fid),
            )
            seq_num = (cursor.fetchone()["cnt"] or 0) + 1
            new_stored_name = build_standardized_filename(target_doc_type, rec["vessel_name"], seq_num, ext)
            new_abs_path = os.path.join(abs_vessel_folder, new_stored_name)
            while os.path.exists(new_abs_path) and os.path.abspath(new_abs_path) != os.path.abspath(old_abs_path):
                seq_num += 1
                new_stored_name = build_standardized_filename(target_doc_type, rec["vessel_name"], seq_num, ext)
                new_abs_path = os.path.join(abs_vessel_folder, new_stored_name)

            if os.path.exists(old_abs_path) and os.path.abspath(old_abs_path) != os.path.abspath(new_abs_path):
                shutil.move(old_abs_path, new_abs_path)
        else:
            new_stored_name = rec["stored_filename"]
            new_abs_path = old_abs_path

        cursor.execute(
            """
            UPDATE uploaded_files
            SET stored_filename = ?,
                file_path = ?,
                doc_type = ?,
                is_wrong_vessel = 0,
                error_code = '',
                error_message = '',
                sync_status = ?
            WHERE id = ?
            """,
            (
                new_stored_name,
                new_stored_name,
                target_doc_type,
                "PENDING_SYNC" if not rec.get("gdrive_file_id") and is_bridge_enabled() else rec.get("sync_status", "LOCAL"),
                fid,
            ),
        )
        cursor.execute(
            "UPDATE checklist_items SET is_manual_verified = 1 WHERE vessel_id = ? AND doc_type = ?",
            (vid, target_doc_type),
        )
        conn.commit()
        conn.close()

        confirmed_count += 1
        background_tasks.add_task(
            bg_rename_doc_and_sync,
            fid,
            vid,
            rec.get("gdrive_file_id", ""),
            new_stored_name,
            folder_meta["folder_name"],
            new_abs_path,
            _mime_for_ext(ext),
            False,
        )

    updated_vessel = None
    for vid in affected_vessel_ids:
        updated_vessel = recalculate_vessel_status(vid)

    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="UPLOAD_FILE",
        details=f"Xác nhận hàng loạt {confirmed_count} chứng từ",
    )

    return {
        "message": f"Đã xác nhận ✓ Đạt cho {confirmed_count} chứng từ đã chọn.",
        "vessel": updated_vessel,
    }


# ==================== 2-WAY GOOGLE DRIVE FOLDER SYNC ENDPOINT ====================

@app.post("/api/vessels/{vessel_id}/sync-drive")
def sync_vessel_with_gdrive_two_way(
    vessel_id: int,
    background_tasks: BackgroundTasks,
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS", "VIEWER")),
):
    """
    Đồng bộ 2 chiều giữa thư mục tàu trên Google Drive và Web:
    1. Quét toàn bộ file trong thư mục tàu trên Google Drive.
    2. Nếu phát hiện file mới do người dùng bỏ trực tiếp lên Drive (ví dụ scan_001.pdf):
       - Tải về đọc nội dung bằng parser_engine.py.
       - Nếu hợp lệ: Đổi tên trực tiếp file đó trên Google Drive thành [{LOAI_GIAY_TO}]_{Ten_Tau}.{ext}
         (không tạo thành 2 file trùng lặp) và cập nhật vào Checklist trên Web.
       - Nếu sai tên tàu hoặc chưa phân loại: Đưa vào danh sách cảnh báo trên Web để User xác nhận.
    3. Nếu có file hợp lệ trên Web chưa có trên Drive: Tự động đẩy lên Drive.
    """
    vessel = get_vessel_detail(vessel_id)
    if not vessel:
        raise HTTPException(status_code=404, detail="Không tìm thấy tàu.")

    if not is_bridge_enabled():
        return {
            "message": "Hệ thống đang chạy ở chế độ Local Storage (chưa cấu hình Google Drive Cloud).",
            "vessel": vessel,
            "new_from_drive": 0,
        }

    folder_meta = get_vessel_folder_meta_internal(vessel_id)
    abs_vessel_folder = get_vessel_storage_abs_path(vessel_id)

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM uploaded_files WHERE vessel_id = ?", (vessel_id,))
    local_rows = [dict(r) for r in cursor.fetchall()]
    conn.close()

    known_drive_ids = [r["gdrive_file_id"] for r in local_rows if r.get("gdrive_file_id")]
    known_stored_names = {r["stored_filename"]: r for r in local_rows}

    drive_res = fetch_vessel_folder_from_gdrive(
        gdrive_folder_id=folder_meta["gdrive_folder_id"],
        vessel_folder_name=folder_meta["folder_name"],
        known_file_ids=known_drive_ids,
    )
    if not drive_res.get("ok"):
        raise HTTPException(
            status_code=400,
            detail=f"Không thể quét thư mục Google Drive: {drive_res.get('error')}",
        )

    # Cập nhật folderId & folderUrl nếu chưa lưu
    if drive_res.get("folderId"):
        conn = get_connection()
        conn.execute(
            "UPDATE vessels SET gdrive_folder_id = ?, gdrive_folder_url = ? WHERE id = ?",
            (drive_res.get("folderId", ""), drive_res.get("folderUrl", ""), vessel_id),
        )
        conn.commit()
        conn.close()

    drive_files = drive_res.get("files") or []
    drive_file_ids_set = {df["fileId"] for df in drive_files if df.get("fileId")}

    imported_count = 0
    renamed_on_drive_count = 0

    for df in drive_files:
        df_id = df.get("fileId") or ""
        df_name = df.get("fileName") or ""
        df_view_url = df.get("viewUrl") or ""
        ext = os.path.splitext(df_name)[1].lower()

        if ext not in (".pdf", ".xlsx", ".docx"):
            continue

        # Nếu fileId đã có trong DB -> bỏ qua
        if df_id in known_drive_ids:
            continue

        # Nếu tên file trùng với stored_filename đã có ở local nhưng chưa gắn gdrive_file_id -> liên kết ID
        if df_name in known_stored_names and not known_stored_names[df_name].get("gdrive_file_id"):
            conn = get_connection()
            conn.execute(
                "UPDATE uploaded_files SET gdrive_file_id = ?, gdrive_view_url = ?, sync_status = 'SYNCED' WHERE id = ?",
                (df_id, df_view_url, known_stored_names[df_name]["id"]),
            )
            conn.commit()
            conn.close()
            known_drive_ids.append(df_id)
            continue

        # Đây là file mới do người dùng tải trực tiếp lên Google Drive!
        raw_b64 = df.get("fileBase64")
        temp_path = os.path.join(abs_vessel_folder, f"__gdrive_{datetime.now().strftime('%H%M%S%f')}_{sanitize_filename(df_name)}")
        if raw_b64:
            try:
                with open(temp_path, "wb") as f_out:
                    f_out.write(base64.b64decode(raw_b64))
            except Exception:
                continue
        else:
            if not download_doc_from_gdrive_if_missing(df_id, temp_path):
                continue

        res = _process_and_store_single_file(
            vessel=vessel,
            temp_file_path=temp_path,
            original_filename=df_name,
            uploaded_by=current_user["username"],
            replace_existing=False,
            preset_gdrive_file_id=df_id,
            preset_gdrive_view_url=df_view_url,
        )
        imported_count += 1

        # Nếu file hợp lệ và tên trên Drive chưa chuẩn -> đổi tên trực tiếp ngay trên Google Drive!
        if not res["is_warning"] and df_name != res["stored_filename"]:
            rn = rename_doc_on_gdrive_sync(df_id, res["stored_filename"], replace_existing=False)
            if rn.get("ok") and rn.get("viewUrl"):
                conn = get_connection()
                conn.execute(
                    "UPDATE uploaded_files SET gdrive_view_url = ?, sync_status = 'SYNCED' WHERE id = ?",
                    (rn["viewUrl"], res["id"]),
                )
                conn.commit()
                conn.close()
                renamed_on_drive_count += 1

    # Nếu có file hợp lệ trên Web chưa được đẩy lên Drive -> đẩy lên Drive
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT * FROM uploaded_files WHERE vessel_id = ? AND is_wrong_vessel = 0 AND doc_type != 'UNCLASSIFIED'",
        (vessel_id,),
    )
    valid_local_files = [dict(r) for r in cursor.fetchall()]
    conn.close()

    for lf in valid_local_files:
        if not lf.get("gdrive_file_id"):
            local_p = os.path.join(abs_vessel_folder, os.path.basename(lf["stored_filename"]))
            if os.path.exists(local_p):
                background_tasks.add_task(
                    bg_upload_doc_and_sync,
                    lf["id"],
                    vessel_id,
                    folder_meta["folder_name"],
                    local_p,
                    lf["stored_filename"],
                    _mime_for_ext(lf.get("file_ext") or ".pdf"),
                )

    updated_vessel = recalculate_vessel_status(vessel_id)
    background_tasks.add_task(sync_db_to_gdrive)

    if imported_count > 0:
        msg = f"Đã đồng bộ {imported_count} file mới từ Google Drive và tự động chuẩn hóa tên file!"
    else:
        msg = "Đã đồng bộ 2 chiều với Google Drive (Hồ sơ đã khớp)."

    return {
        "message": msg,
        "new_from_drive": imported_count,
        "renamed_on_drive": renamed_on_drive_count,
        "vessel": updated_vessel,
    }


@app.post("/api/demo/reset")
def reset_demo_environment(
    background_tasks: BackgroundTasks = BackgroundTasks(),
    current_user: Dict[str, Any] = Depends(require_roles("ADMIN", "OPS")),
):
    seed_initial_vessels(reset_existing=True)
    generate_all_sample_files()
    vessels = get_all_vessels()
    log_audit_event(
        username=current_user["username"],
        role=current_user["role"],
        action="CREATE_VESSEL",
        details="Khôi phục dữ liệu tàu mẫu mặc định (MV BINH HOA & MV HAI AN)",
    )
    background_tasks.add_task(sync_db_to_gdrive)
    return {
        "message": "Đã khôi phục dữ liệu mẫu mặc định.",
        "vessels": vessels,
    }


if __name__ == "__main__":
    import socket
    import uvicorn

    env_port = os.environ.get("PORT")
    if env_port and env_port.isdigit():
        chosen_port = int(env_port)
        host_addr = "0.0.0.0"
    else:
        host_addr = "127.0.0.1"
        preferred_ports = [8000, 8080, 8001, 8888, 5000]
        chosen_port = 8080
        for p in preferred_ports:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                try:
                    s.bind(("127.0.0.1", p))
                    chosen_port = p
                    break
                except OSError:
                    continue

    print(f"Starting PortCall Docs Enterprise MVP at: http://{host_addr}:{chosen_port}")
    uvicorn.run(app, host=host_addr, port=chosen_port)
