import os
import re
import sys
import warnings
import unicodedata
from typing import Dict, Any, List, Tuple

warnings.filterwarnings("ignore", category=FutureWarning)
import fitz  # PyMuPDF
import openpyxl
import docx  # python-docx


# Danh sách từ khóa mỏ neo (Anchor Keywords) theo đúng đặc tả nghiệp vụ
ANCHOR_KEYWORDS: List[Tuple[str, List[str]]] = [
    ("VGM", ["VERIFIED GROSS MASS", "VGM DECLARATION"]),
    ("MANIFEST", ["CARGO MANIFEST", "LƯỢC KHAI HÀNG HÓA", "LUOC KHAI HANG HOA"]),
    ("BL", ["BILL OF LADING", "OCEAN BILL"]),
    ("CREW_LIST", ["CREW LIST", "IMO FAL FORM 5"]),
    ("PORT_CLEARANCE", ["PORT CLEARANCE", "GIẤY PHÉP RỜI CẢNG", "GIAY PHEP ROI CANG"]),
    ("STOWAGE_PLAN", ["STOWAGE PLAN", "SƠ ĐỒ XẾP HÀNG", "SO DO XEP HANG"]),
    ("MATES_RECEIPT", ["MATE'S RECEIPT", "MATES RECEIPT", "BIÊN LAI THUYỀN PHÓ", "BIEN LAI THUYEN PHO"]),
]

# Từ khóa dự phòng (Fallback) khi quét tên file gốc nếu nội dung không khớp
FILENAME_FALLBACK_PATTERNS: List[Tuple[str, str]] = [
    ("VGM", r"(?:VERIFIED[\s_\-]*GROSS[\s_\-]*MASS|VGM[\s_\-]*DECLARATION|\bVGM\b)"),
    ("MANIFEST", r"(?:CARGO[\s_\-]*MANIFEST|LƯỢC[\s_\-]*KHAI|LUOC[\s_\-]*KHAI|\bMANIFEST\b)"),
    ("BL", r"(?:BILL[\s_\-]*OF[\s_\-]*LADING|OCEAN[\s_\-]*BILL|\bBL\b|\bB_L\b)"),
    ("CREW_LIST", r"(?:CREW[\s_\-]*LIST|IMO[\s_\-]*FAL[\s_\-]*FORM[\s_\-]*5|\bCREW\b)"),
    ("PORT_CLEARANCE", r"(?:PORT[\s_\-]*CLEARANCE|GIẤY[\s_\-]*PHÉP[\s_\-]*RỜI[\s_\-]*CẢNG|GIAY[\s_\-]*PHEP[\s_\-]*ROI[\s_\-]*CANG|\bCLEARANCE\b)"),
    ("STOWAGE_PLAN", r"(?:STOWAGE[\s_\-]*PLAN|\bSTOWAGE\b)"),
    ("MATES_RECEIPT", r"(?:MATE'?S?[\s_\-]*RECEIPT|\bMATES\b)"),
]


def sanitize_token(text: str) -> str:
    """Chuẩn hóa chuỗi để đặt tên file chuẩn an toàn."""
    if not text:
        return "UNKNOWN"
    cleaned = re.sub(r"[^\w\-]", "_", text.strip())
    cleaned = re.sub(r"_+", "_", cleaned).strip("_")
    return cleaned.upper() if cleaned else "UNKNOWN"


def build_standardized_filename(doc_type: str, vessel_name: str, seq_num: int = 1, ext: str = ".pdf") -> str:
    """
    Quy tắc đặt tên file chuẩn rút gọn (không kèm mã chuyến V01 và không gắn _01 mặc định):
    - File thứ nhất (seq_num == 1): [{LOAI_GIAY_TO}]_{Ten_Tau}.{ext}
      Ví dụ: [MANIFEST]_MV_BINH_HOA.pdf
    - File thứ 2 trở đi cùng loại (seq_num >= 2): [{LOAI_GIAY_TO}]_{Ten_Tau}_({seq_num}).{ext}
      Ví dụ: [VGM]_MV_BINH_HOA_(2).pdf
    """
    clean_ext = ext if ext.startswith(".") else f".{ext}"
    clean_vessel = sanitize_token(vessel_name)
    clean_type = sanitize_token(doc_type)
    if seq_num <= 1:
        return f"[{clean_type}]_{clean_vessel}{clean_ext.lower()}"
    return f"[{clean_type}]_{clean_vessel}_({seq_num}){clean_ext.lower()}"


def strip_accents(text: str) -> str:
    """Loại bỏ dấu tiếng Việt và chuẩn hóa chuỗi để đối chiếu không phân biệt hoa thường/dấu."""
    if not text:
        return ""
    text = text.replace("Đ", "D").replace("đ", "d")
    normalized = unicodedata.normalize("NFD", text)
    without_accents = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", without_accents).strip().upper()


def normalize_vessel_name_for_compare(name: str) -> str:
    """Chuẩn hóa tên tàu để so sánh (không phân biệt hoa/thường, dấu, khoảng trắng thừa)."""
    norm = strip_accents(name)
    norm = re.sub(r"[^A-Z0-9\s]", " ", norm)
    norm = re.sub(r"\s+", " ", norm).strip()
    return norm


def is_same_vessel(vessel_in_file: str, current_vessel: str) -> bool:
    """
    So sánh tên tàu trong file và tên tàu hiện tại.
    Trả về True nếu khớp (kể cả khi một bên có/không có tiền tố MV / M/V), False nếu khác tàu.
    """
    v1 = normalize_vessel_name_for_compare(vessel_in_file)
    v2 = normalize_vessel_name_for_compare(current_vessel)
    if not v1 or not v2:
        return True
    if v1 == v2:
        return True
    v1_core = re.sub(r"^(?:MV|M\s+V)\s+", "", v1).strip()
    v2_core = re.sub(r"^(?:MV|M\s+V)\s+", "", v2).strip()
    return bool(v1_core and v2_core and v1_core == v2_core)


def normalize_extracted_text(text: str) -> str:
    """Chuẩn hóa các ký tự khoảng trắng đặc biệt (\\xa0) và dấu gạch ngang Unicode (\\xad, –...) khi rút trích."""
    if not text:
        return ""
    text = text.replace("\xa0", " ").replace("\u200b", "")
    text = re.sub(r"[\xad\u2010\u2011\u2012\u2013\u2014\u2212]", "-", text)
    return text


def extract_text_from_pdf(file_path: str) -> Tuple[str, str]:
    """
    Đọc text từ file PDF bằng PyMuPDF (fitz).
    Trả về (first_page_text, full_text).
    """
    first_page_text = ""
    all_pages = []
    try:
        doc = fitz.open(file_path)
        for idx, page in enumerate(doc):
            page_txt = normalize_extracted_text(page.get_text("text") or "")
            if idx == 0:
                first_page_text = page_txt
            all_pages.append(page_txt)
        doc.close()
    except Exception as exc:
        print(f"[ParserEngine] Lỗi đọc PDF {file_path}: {exc}")
    full_text = "\n".join(all_pages)
    return first_page_text, full_text


def extract_text_from_excel(file_path: str) -> Tuple[str, str]:
    """
    Đọc text từ file Excel (.xlsx) bằng openpyxl.
    Trả về (first_sheet_text, full_text).
    """
    first_sheet_text = ""
    all_sheets = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            lines = []
            for row in ws.iter_rows(values_only=True):
                row_vals = [
                    normalize_extracted_text(str(cell).strip())
                    for cell in row
                    if cell is not None and str(cell).strip() != ""
                ]
                if row_vals:
                    lines.append(" ".join(row_vals))
            sheet_txt = "\n".join(lines)
            if idx == 0:
                first_sheet_text = sheet_txt
            all_sheets.append(sheet_txt)
        wb.close()
    except Exception as exc:
        print(f"[ParserEngine] Lỗi đọc Excel {file_path}: {exc}")
    full_text = "\n".join(all_sheets)
    return first_sheet_text, full_text


def extract_text_from_docx(file_path: str) -> Tuple[str, str]:
    """
    Đọc text từ file Word (.docx) bằng python-docx (bao gồm cả đoạn văn và bảng biểu).
    Giữ nguyên định dạng gốc của file tải lên, không convert sang PDF/Word khác.
    Trả về (first_section_text, full_text).
    """
    para_lines = []
    table_lines = []
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            txt = normalize_extracted_text(p.text.strip())
            if txt:
                para_lines.append(txt)

        for table in doc.tables:
            for row in table.rows:
                row_cells = []
                last_val = None
                for cell in row.cells:
                    c_txt = normalize_extracted_text(cell.text.strip().replace("\n", " "))
                    if c_txt and c_txt != last_val:
                        row_cells.append(c_txt)
                        last_val = c_txt
                if row_cells:
                    table_lines.append(" | ".join(row_cells))
    except Exception as exc:
        print(f"[ParserEngine] Lỗi đọc Word (.docx) {file_path}: {exc}")

    all_lines = para_lines + table_lines
    full_text = "\n".join(all_lines)
    first_page_text = "\n".join(all_lines[:40])
    return first_page_text, full_text


def classify_document(first_page_text: str, original_filename: str) -> Tuple[str, str, str]:
    """
    Phân loại chứng từ dựa trên:
    1. Quét từ khóa mỏ neo (Anchor Keywords) ở trang đầu tiên.
    2. Fallback: Quét tên file gốc nếu nội dung không khớp.
    3. Nếu vẫn không nhận diện được -> UNCLASSIFIED.
    Trả về: (doc_type, detected_method, detected_keyword)
    """
    if first_page_text and first_page_text.strip():
        upper_text = first_page_text.upper()
        unaccented_text = strip_accents(first_page_text)
        for doc_type, keywords in ANCHOR_KEYWORDS:
            for kw in keywords:
                kw_upper = kw.upper()
                if kw_upper in upper_text or strip_accents(kw_upper) in unaccented_text:
                    return doc_type, "CONTENT_REGEX", kw

    base_name = os.path.splitext(original_filename or "")[0]
    upper_name = base_name.upper()
    unaccented_name = strip_accents(base_name)

    for doc_type, keywords in ANCHOR_KEYWORDS:
        for kw in keywords:
            kw_norm = kw.upper().replace(" ", "[\\s_\\-]*")
            if re.search(kw_norm, upper_name, re.IGNORECASE) or re.search(
                strip_accents(kw).replace(" ", "[\\s_\\-]*"), unaccented_name, re.IGNORECASE
            ):
                return doc_type, "FILENAME_FALLBACK", f"Filename: {kw}"

    for doc_type, pattern in FILENAME_FALLBACK_PATTERNS:
        m = re.search(pattern, upper_name, re.IGNORECASE) or re.search(pattern, unaccented_name, re.IGNORECASE)
        if m:
            return doc_type, "FILENAME_FALLBACK", f"Filename: {m.group(0)}"

    return "UNCLASSIFIED", "UNCLASSIFIED", ""


def extract_vessel_name_from_text(text: str) -> str:
    """
    Dùng Regex tìm chuỗi đứng sau 'Vessel Name:', 'Vessel:' hoặc 'M/V:' trong nội dung file.
    """
    if not text:
        return ""
    pattern = re.compile(
        r"(?:Vessel\s*Name|Vessel|M/V)\s*[:\-]\s*([^\r\n\|\;\,]+)",
        re.IGNORECASE,
    )
    match = pattern.search(text)
    if not match:
        return ""
    raw_vessel = match.group(1).strip()
    raw_vessel = re.split(
        r"\s+(?:Voyage\s*No|Voyage|VOY|ETA|Port\s+of|Call\s+Sign|IMO\s+No|Flag|Date)\b",
        raw_vessel,
        flags=re.IGNORECASE,
    )[0].strip()
    raw_vessel = raw_vessel.strip(" .,-_|")
    return raw_vessel


def extract_shipment_counts_and_codes(text: str, doc_type: str) -> Tuple[int, List[str]]:
    """
    Đối chiếu số lượng lô hàng (B/L hoặc Container):
    - Tìm 'Total B/Ls: (\\d+)' (đặc biệt trong MANIFEST)
    - Đếm các mã vận đơn chuẩn 'BL-\\d+'
    - Hoặc đếm số mã Container chuẩn ISO '[A-Z]{4}\\d{7}'
    Trả về: (extracted_count, unique_codes_list)
    """
    if not text:
        return 0, []

    bl_matches = re.findall(r"\bBL-\d+\b", text, flags=re.IGNORECASE)
    unique_bls = []
    seen_bls = set()
    for code in bl_matches:
        c_up = code.upper()
        if c_up not in seen_bls:
            seen_bls.add(c_up)
            unique_bls.append(c_up)

    cont_matches = re.findall(r"\b[A-Z]{4}\d{7}\b", text)
    unique_conts = []
    seen_conts = set()
    for code in cont_matches:
        c_up = code.upper()
        if c_up not in seen_conts:
            seen_conts.add(c_up)
            unique_conts.append(c_up)

    codes = unique_bls if len(unique_bls) > 0 else unique_conts

    total_bl_match = re.search(r"Total\s+B/Ls?\s*[:\-]\s*(\d+)", text, flags=re.IGNORECASE)
    declared_total = int(total_bl_match.group(1)) if total_bl_match else 0

    if doc_type == "MANIFEST" and declared_total > 0:
        count = max(declared_total, len(codes))
    else:
        count = len(codes) if len(codes) > 0 else declared_total

    return count, codes


def parse_and_validate_file(
    file_path: str,
    original_filename: str,
    current_vessel_name: str,
    forced_doc_type: str = "",
) -> Dict[str, Any]:
    """
    Hàm tổng hợp xử lý 1 file (.pdf, .xlsx, .docx):
    1. Trích xuất văn bản (PyMuPDF cho PDF, openpyxl cho Excel, python-docx cho Word).
    2. Tự động phân loại chứng từ bằng Anchor Keywords (hoặc dùng forced_doc_type khi chọn 1-click).
    3. Kiểm tra chéo tên tàu (Cross-validation WRONG_VESSEL).
    4. Đếm số lượng B/L hoặc Container phục vụ đối chiếu với Manifest.
    """
    ext = os.path.splitext(original_filename or file_path)[1].lower()
    if ext == ".pdf":
        first_page_text, full_text = extract_text_from_pdf(file_path)
    elif ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        first_page_text, full_text = extract_text_from_excel(file_path)
    elif ext == ".docx":
        first_page_text, full_text = extract_text_from_docx(file_path)
    else:
        first_page_text, full_text = "", ""

    if forced_doc_type and forced_doc_type != "UNCLASSIFIED":
        doc_type = forced_doc_type
        detected_method = "MANUAL_1CLICK"
        detected_keyword = f"1-Click: {forced_doc_type}"
    else:
        doc_type, detected_method, detected_keyword = classify_document(first_page_text, original_filename)

    detected_vessel = extract_vessel_name_from_text(first_page_text)
    if not detected_vessel and full_text:
        detected_vessel = extract_vessel_name_from_text(full_text)

    is_wrong_vessel = 0
    error_code = ""
    error_message = ""

    if detected_vessel and not is_same_vessel(detected_vessel, current_vessel_name):
        is_wrong_vessel = 1
        error_code = "WRONG_VESSEL"
        error_message = (
            f"Cảnh báo: Nội dung file ghi tàu {detected_vessel}, "
            f"không khớp với tàu {current_vessel_name}!"
        )
    elif doc_type == "UNCLASSIFIED":
        error_code = "UNCLASSIFIED"
        error_message = "Không nhận diện được loại chứng từ tự động. Vui lòng chọn loại giấy tờ bằng nút Dropdown (1-click)."

    extracted_count, extracted_codes = extract_shipment_counts_and_codes(full_text or first_page_text, doc_type)

    preview_source = (first_page_text or full_text or "").strip()
    preview_text = re.sub(r"\n{3,}", "\n\n", preview_source)[:600]

    return {
        "doc_type": doc_type,
        "detected_method": detected_method,
        "detected_keyword": detected_keyword,
        "detected_vessel_name": detected_vessel,
        "is_wrong_vessel": is_wrong_vessel,
        "error_code": error_code,
        "error_message": error_message,
        "extracted_count": extracted_count,
        "extracted_codes": extracted_codes,
        "extracted_text_preview": preview_text,
    }
